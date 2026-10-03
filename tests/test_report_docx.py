from io import BytesIO
import unittest

from docx import Document
from docx.shared import Mm

from web.report_docx import (
    build_report_docx,
    validate_report_draft,
)


def draft_fixture():
    return {
        "schema_version":
            "mip-report-draft/1",
        "items": [
            {
                "type": "signal",
                "section_id": "context",
                "title":
                    "Тестовий сигнал",
                "factual_summary":
                    "Короткий фактичний виклад.",
                "analyst_assessment":
                    "Оцінка аналітика.",
                "source_count": 5,
                "occurrence_count": 8,
            },
            {
                "type": "material",
                "section_id": "enemy",
                "title":
                    "Тестовий матеріал",
                "factual_summary":
                    "Факт з матеріалу.",
                "analyst_assessment": "",
                "source_name":
                    "Тестове джерело",
                "source_url":
                    "https://example.com/item",
            },
        ],
    }


class ReportDraftValidationTests(
    unittest.TestCase
):
    def test_unassigned_item_is_rejected(self):
        payload = draft_fixture()
        payload["items"][0][
            "section_id"
        ] = None

        with self.assertRaises(
            ValueError
        ):
            validate_report_draft(
                payload
            )

    def test_unknown_section_is_rejected(self):
        payload = draft_fixture()
        payload["items"][0][
            "section_id"
        ] = "unknown"

        with self.assertRaises(
            ValueError
        ):
            validate_report_draft(
                payload
            )


class ReportDocxTests(unittest.TestCase):
    def build(self):
        payload = build_report_docx(
            draft_fixture(),
            report_date="02.10.2026",
        )

        self.assertTrue(
            payload.startswith(b"PK")
        )

        return Document(
            BytesIO(payload)
        )

    def test_document_has_a4_layout(self):
        document = self.build()
        section = document.sections[0]

        self.assertAlmostEqual(
            section.page_width.mm,
            Mm(210).mm,
            places=1,
        )
        self.assertAlmostEqual(
            section.page_height.mm,
            Mm(297).mm,
            places=1,
        )

    def test_document_contains_numbered_sections_and_items(self):
        document = self.build()

        texts = [
            paragraph.text
            for paragraph
            in document.paragraphs
        ]

        self.assertIn(
            "АНАЛІТИЧНИЙ ЗВІТ",
            texts,
        )
        self.assertIn(
            "1. Міжнародний та "
            "загальнодержавний контекст",
            texts,
        )
        self.assertIn(
            "1.1. Тестовий сигнал",
            texts,
        )
        self.assertIn(
            "2. Інформація з ворожого "
            "інформаційного простору",
            texts,
        )
        self.assertIn(
            "2.1. Тестовий матеріал",
            texts,
        )

    def test_assessment_is_written_only_when_present(self):
        document = self.build()

        paragraphs = [
            paragraph.text
            for paragraph
            in document.paragraphs
        ]

        self.assertEqual(
            paragraphs.count(
                "Оцінка аналітика"
            ),
            1,
        )

        self.assertIn(
            "Оцінка аналітика.",
            paragraphs,
        )

    def test_source_and_signal_meta_are_preserved(self):
        document = self.build()

        texts = "\n".join(
            paragraph.text
            for paragraph
            in document.paragraphs
        )

        self.assertIn(
            "Сигнал МІП: 5 джерел "
            "· 8 публікацій",
            texts,
        )
        self.assertIn(
            "Тестове джерело · "
            "Відкрити публікацію ↗",
            texts,
        )

        urls = [
            hyperlink.url
            for paragraph
            in document.paragraphs
            for hyperlink
            in paragraph.hyperlinks
        ]

        self.assertIn(
            "https://example.com/item",
            urls,
        )


if __name__ == "__main__":
    unittest.main()



class ReportSummaryOverrideValidationTests(
    unittest.TestCase
):
    def test_signal_override_empty_string_is_preserved(self):
        from web.report_docx import (
            validate_report_draft,
        )

        result = validate_report_draft({
            "schema_version":
                "mip-report-draft/1",
            "items": [
                {
                    "type": "signal",
                    "section_id": "context",
                    "title": "Тестовий сигнал",
                    "factual_summary":
                        "Auto snapshot",
                    "factual_summary_override":
                        "",
                    "analyst_assessment": "",
                    "source_name": "",
                    "source_url": "",
                    "source_count": 1,
                    "occurrence_count": 1,
                }
            ],
        })

        self.assertEqual(
            result["items"][0][
                "factual_summary_override"
            ],
            "",
        )

    def test_signal_without_override_keeps_none(self):
        from web.report_docx import (
            validate_report_draft,
        )

        result = validate_report_draft({
            "schema_version":
                "mip-report-draft/1",
            "items": [
                {
                    "type": "signal",
                    "section_id": "context",
                    "title": "Тестовий сигнал",
                    "factual_summary":
                        "Auto snapshot",
                    "analyst_assessment": "",
                    "source_name": "",
                    "source_url": "",
                    "source_count": 1,
                    "occurrence_count": 1,
                }
            ],
        })

        self.assertIsNone(
            result["items"][0][
                "factual_summary_override"
            ]
        )



class ReportDocxMetadataTests(
    unittest.TestCase
):
    def test_generated_document_has_clean_metadata(self):
        from datetime import (
            datetime,
        )
        from io import BytesIO
        from xml.etree import (
            ElementTree,
        )
        from zipfile import ZipFile
        from zoneinfo import ZoneInfo

        from docx import Document

        from web.report_docx import (
            build_report_docx,
        )

        draft = {
            "schema_version":
                "mip-report-draft/1",
            "items": [
                {
                    "type": "material",
                    "section_id":
                        "context",
                    "title":
                        "Тестовий матеріал",
                    "factual_summary":
                        "Фактичний виклад.",
                    "analyst_assessment":
                        "",
                }
            ],
        }

        generated = (
            build_report_docx(
                draft,
                report_date=
                    "03.10.2026",
            )
        )

        document = Document(
            BytesIO(generated)
        )

        properties = (
            document.core_properties
        )

        self.assertEqual(
            properties.author,
            "",
        )
        self.assertEqual(
            properties.last_modified_by,
            "",
        )
        self.assertEqual(
            properties.comments,
            "",
        )

        now_utc = (
            datetime.now(
                ZoneInfo("UTC")
            )
        )

        self.assertIsNotNone(
            properties.created
        )
        self.assertIsNotNone(
            properties.modified
        )

        self.assertLess(
            abs(
                (
                    now_utc
                    - properties.created
                ).total_seconds()
            ),
            60,
        )

        with ZipFile(
            BytesIO(generated),
            "r",
        ) as archive:
            root = (
                ElementTree.fromstring(
                    archive.read(
                        "docProps/app.xml"
                    )
                )
            )

        namespace = (
            "http://schemas.openxmlformats.org/"
            "officeDocument/2006/"
            "extended-properties"
        )

        def value(name):
            element = root.find(
                "{"
                + namespace
                + "}"
                + name
            )

            if element is None:
                return None

            return element.text

        self.assertEqual(
            value("Application"),
            "МІП Report Builder",
        )

        for name in (
            "Template",
            "TotalTime",
            "Pages",
            "Words",
            "Characters",
            "CharactersWithSpaces",
            "Lines",
            "Paragraphs",
            "AppVersion",
        ):
            self.assertIsNone(
                value(name),
                name,
            )
