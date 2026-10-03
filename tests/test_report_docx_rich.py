from io import BytesIO
import unittest

from docx import Document

from web.report_docx import (
    build_report_docx,
)


def rich_signal(
    *,
    with_theses=True,
):
    item = {
        "type": "signal",
        "section_id": "context",
        "title":
            "Тестовий сигнал",
        "factual_summary": "",
        "analyst_assessment":
            "Оцінка аналітика.",
        "source_name":
            "Джерело 1",
        "source_url":
            "https://example.com/post",
        "source_count": 4,
        "occurrence_count": 7,
        "signal_summary":
            (
                "Стислий зміст сигналу."
                if with_theses
                else ""
            ),
        "signal_theses":
            (
                [
                    "Перша теза.",
                    "Друга теза.",
                ]
                if with_theses
                else []
            ),
        "example_publication": {
            "source_name":
                "Джерело 1",
            "source_type":
                "telegram",
            "source_group":
                "ru_space",
            "title":
                "Приклад",
            "text":
                "Оригінальний текст поста.",
            "url":
                "https://example.com/post",
            "observed_at":
                "2026-10-02T10:00:00+00:00",
        },
    }

    return {
        "schema_version":
            "mip-report-draft/1",
        "items": [item],
    }


class RichReportDocxTests(
    unittest.TestCase
):
    def build(
        self,
        *,
        with_theses=True,
    ):
        payload = build_report_docx(
            rich_signal(
                with_theses=
                    with_theses,
            ),
            report_date="02.10.2026",
            resolved=True,
        )

        return Document(
            BytesIO(payload)
        )

    def test_signal_renders_summary_theses_and_example(self):
        document = self.build()

        text = "\n".join(
            paragraph.text
            for paragraph
            in document.paragraphs
        )

        self.assertIn(
            "Коротко",
            text,
        )
        self.assertIn(
            "Стислий зміст сигналу.",
            text,
        )
        self.assertIn(
            "Основні тези",
            text,
        )
        self.assertIn(
            "1) Перша теза.",
            text,
        )
        self.assertIn(
            "Приклад публікації",
            text,
        )
        self.assertIn(
            "Оригінальний текст поста.",
            text,
        )

    def test_signal_without_theses_renders_only_example(self):
        document = self.build(
            with_theses=False,
        )

        text = "\n".join(
            paragraph.text
            for paragraph
            in document.paragraphs
        )

        self.assertNotIn(
            "Коротко",
            text,
        )
        self.assertNotIn(
            "Основні тези",
            text,
        )
        self.assertIn(
            "Приклад публікації",
            text,
        )

    def test_external_links_are_real_hyperlinks(self):
        document = self.build()

        urls = [
            hyperlink.url
            for paragraph
            in document.paragraphs
            for hyperlink
            in paragraph.hyperlinks
        ]

        self.assertIn(
            "https://example.com/post",
            urls,
        )

        self.assertGreaterEqual(
            urls.count(
                "https://example.com/post"
            ),
            2,
        )

    def test_item_has_more_vertical_spacing(self):
        document = self.build()

        heading = next(
            paragraph
            for paragraph
            in document.paragraphs
            if paragraph.text.startswith(
                "1.1. "
            )
        )

        self.assertGreaterEqual(
            heading.paragraph_format
            .space_before.pt,
            12,
        )


if __name__ == "__main__":
    unittest.main()
