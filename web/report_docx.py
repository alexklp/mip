"""Формування редагованого DOCX з локальної чернетки звіту МІП."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from zipfile import ZipFile
from typing import Any
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo
from xml.etree import ElementTree

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt


REPORT_SCHEMA_VERSION = "mip-report-draft/1"

MAX_ITEMS = 100
MAX_TITLE_CHARS = 600
MAX_FACTUAL_CHARS = 12000
MAX_ASSESSMENT_CHARS = 8000
MAX_SOURCE_NAME_CHARS = 300
MAX_SOURCE_URL_CHARS = 2048

SECTION_ORDER = (
    (
        "context",
        "Міжнародний та загальнодержавний контекст",
    ),
    (
        "forces",
        "Інформація про діяльність ЗСУ та ДШВ",
    ),
    (
        "negative",
        "Негативна інформація, пов’язана із ЗСУ та ДШВ",
    ),
    (
        "enemy",
        "Інформація з ворожого інформаційного простору",
    ),
)

SECTION_IDS = {
    section_id
    for section_id, _ in SECTION_ORDER
}


def validate_report_draft(
    payload: Any,
) -> dict:
    """Перевірити bounded client draft перед DOCX export."""
    if not isinstance(payload, dict):
        raise ValueError(
            "Чернетка звіту має бути JSON object"
        )

    if (
        payload.get("schema_version")
        != REPORT_SCHEMA_VERSION
    ):
        raise ValueError(
            "Непідтримувана версія чернетки звіту"
        )

    items = payload.get("items")

    if (
        not isinstance(items, list)
        or not items
        or len(items) > MAX_ITEMS
    ):
        raise ValueError(
            "Звіт має містити від 1 до 100 елементів"
        )

    normalized_items = []

    for index, item in enumerate(items):
        if not isinstance(item, dict):
            raise ValueError(
                f"Елемент {index + 1}: некоректний формат"
            )

        item_type = item.get("type")

        if item_type not in {
            "signal",
            "material",
        }:
            raise ValueError(
                f"Елемент {index + 1}: некоректний тип"
            )

        section_id = item.get(
            "section_id"
        )

        if section_id not in SECTION_IDS:
            raise ValueError(
                "Усі елементи потрібно розподілити "
                "за розділами перед експортом"
            )

        title = _required_text(
            item.get("title"),
            field=(
                f"Елемент {index + 1}: заголовок"
            ),
            max_chars=MAX_TITLE_CHARS,
        )

        factual = _optional_text(
            item.get("factual_summary"),
            field=(
                f"Елемент {index + 1}: фактичний виклад"
            ),
            max_chars=MAX_FACTUAL_CHARS,
        )

        factual_override = None

        if (
            item_type == "signal"
            and "factual_summary_override" in item
        ):
            raw_override = item.get(
                "factual_summary_override"
            )

            if raw_override is not None:
                if not isinstance(
                    raw_override,
                    str,
                ):
                    raise ValueError(
                        f"Елемент {index + 1}: "
                        "ручний фактичний виклад "
                        "має бути текстом"
                    )

                factual_override = (
                    _optional_text(
                        raw_override,
                        field=(
                            f"Елемент {index + 1}: "
                            "ручний фактичний виклад"
                        ),
                        max_chars=
                            MAX_FACTUAL_CHARS,
                    )
                )

                # Порожній рядок має окрему семантику:
                # аналітик свідомо прибрав «Коротко».
                if not raw_override.strip():
                    factual_override = ""

        assessment = _optional_text(
            item.get("analyst_assessment"),
            field=(
                f"Елемент {index + 1}: оцінка аналітика"
            ),
            max_chars=MAX_ASSESSMENT_CHARS,
        )

        source_name = _optional_text(
            item.get("source_name"),
            field=(
                f"Елемент {index + 1}: джерело"
            ),
            max_chars=MAX_SOURCE_NAME_CHARS,
        )

        source_url = _optional_url(
            item.get("source_url"),
            field=(
                f"Елемент {index + 1}: URL джерела"
            ),
        )

        source_count = _optional_nonnegative_int(
            item.get("source_count"),
            field=(
                f"Елемент {index + 1}: source_count"
            ),
        )

        occurrence_count = (
            _optional_nonnegative_int(
                item.get("occurrence_count"),
                field=(
                    f"Елемент {index + 1}: "
                    "occurrence_count"
                ),
            )
        )

        normalized_items.append(
            {
                "type": item_type,
                "section_id": section_id,
                "title": title,
                "factual_summary": factual,
                "factual_summary_override":
                    factual_override,
                "analyst_assessment":
                    assessment,
                "source_name": source_name,
                "source_url": source_url,
                "source_count": source_count,
                "occurrence_count":
                    occurrence_count,
            }
        )

    return {
        "schema_version":
            REPORT_SCHEMA_VERSION,
        "items": normalized_items,
    }



APP_PROPERTIES_NS = (
    "http://schemas.openxmlformats.org/"
    "officeDocument/2006/extended-properties"
)

VT_PROPERTIES_NS = (
    "http://schemas.openxmlformats.org/"
    "officeDocument/2006/docPropsVTypes"
)


def _configure_document_metadata(
    document: Document,
) -> None:
    """Прибрати template provenance та записати реальний час створення."""
    generated_at = (
        datetime.now(
            ZoneInfo("UTC")
        )
        .replace(
            tzinfo=None,
            microsecond=0,
        )
    )

    properties = (
        document.core_properties
    )

    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""
    properties.created = generated_at
    properties.modified = generated_at
    properties.revision = 1


def _sanitize_extended_properties(
    docx_bytes: bytes,
) -> bytes:
    """Прибрати недостовірні extended properties шаблону python-docx."""
    source = BytesIO(docx_bytes)
    target = BytesIO()

    ElementTree.register_namespace(
        "",
        APP_PROPERTIES_NS,
    )
    ElementTree.register_namespace(
        "vt",
        VT_PROPERTIES_NS,
    )

    with ZipFile(source, "r") as input_zip:
        with ZipFile(
            target,
            "w",
        ) as output_zip:
            for info in input_zip.infolist():
                data = input_zip.read(
                    info.filename
                )

                if (
                    info.filename
                    == "docProps/app.xml"
                ):
                    root = (
                        ElementTree.fromstring(
                            data
                        )
                    )

                    application = root.find(
                        (
                            "{"
                            + APP_PROPERTIES_NS
                            + "}Application"
                        )
                    )

                    if application is None:
                        application = (
                            ElementTree.SubElement(
                                root,
                                (
                                    "{"
                                    + APP_PROPERTIES_NS
                                    + "}Application"
                                ),
                            )
                        )

                    application.text = (
                        "МІП Report Builder"
                    )

                    for name in (
                        "Template",
                        "TotalTime",
                        "Pages",
                        "Words",
                        "Characters",
                        "CharactersWithSpaces",
                        "Lines",
                        "Paragraphs",
                        "AppVersion",
                    ):
                        element = root.find(
                            (
                                "{"
                                + APP_PROPERTIES_NS
                                + "}"
                                + name
                            )
                        )

                        if element is not None:
                            root.remove(
                                element
                            )

                    data = (
                        ElementTree.tostring(
                            root,
                            encoding="utf-8",
                            xml_declaration=True,
                        )
                    )

                output_zip.writestr(
                    info,
                    data,
                )

    return target.getvalue()



def build_report_docx(
    draft: dict,
    *,
    report_date: str,
    resolved: bool = False,
) -> bytes:
    """Побудувати редагований A4 DOCX без зміни даних МІП."""
    validated = (
        draft
        if resolved
        else validate_report_draft(
            draft
        )
    )

    document = Document()

    _configure_document_metadata(
        document
    )

    _configure_document(
        document
    )

    title = document.add_paragraph()
    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        "АНАЛІТИЧНИЙ ЗВІТ"
    )
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    _set_run_font(run)

    date_paragraph = (
        document.add_paragraph()
    )
    date_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    date_run = date_paragraph.add_run(
        report_date
    )
    date_run.font.name = (
        "Times New Roman"
    )
    date_run.font.size = Pt(11)
    _set_run_font(date_run)

    section_number = 0

    for (
        section_id,
        section_title,
    ) in SECTION_ORDER:
        section_items = [
            item
            for item in validated["items"]
            if item["section_id"]
            == section_id
        ]

        if not section_items:
            continue

        section_number += 1

        _add_section_heading(
            document,
            number=section_number,
            title=section_title,
        )

        for item_number, item in enumerate(
            section_items,
            start=1,
        ):
            _add_report_item(
                document,
                item=item,
                number=(
                    f"{section_number}."
                    f"{item_number}"
                ),
            )

    output = BytesIO()
    document.save(output)

    return _sanitize_extended_properties(
        output.getvalue()
    )


def _configure_document(
    document: Document,
) -> None:
    section = document.sections[0]

    section.start_type = (
        WD_SECTION.NEW_PAGE
    )
    section.page_width = Mm(210)
    section.page_height = Mm(297)

    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    normal = document.styles["Normal"]
    normal.font.name = (
        "Times New Roman"
    )
    normal.font.size = Pt(12)

    normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman",
    )

    paragraph = (
        normal.paragraph_format
    )
    paragraph.space_after = Pt(6)
    paragraph.line_spacing = 1.15


def _add_section_heading(
    document: Document,
    *,
    number: int,
    title: str,
) -> None:
    paragraph = (
        document.add_paragraph()
    )

    paragraph.paragraph_format.space_before = (
        Pt(16)
    )
    paragraph.paragraph_format.space_after = (
        Pt(8)
    )
    paragraph.paragraph_format.keep_with_next = True

    run = paragraph.add_run(
        f"{number}. {title}"
    )

    run.bold = True
    run.font.name = (
        "Times New Roman"
    )
    run.font.size = Pt(14)

    _set_run_font(run)


def _add_report_item(
    document: Document,
    *,
    item: dict,
    number: str,
) -> None:
    example = item.get(
        "example_publication"
    )

    heading_url = (
        (
            example.get("url")
            if isinstance(
                example,
                dict,
            )
            else ""
        )
        or item.get("source_url")
        or ""
    )

    heading = document.add_paragraph()

    heading.paragraph_format.space_before = (
        Pt(14)
    )
    heading.paragraph_format.space_after = (
        Pt(5)
    )
    heading.paragraph_format.keep_with_next = True

    number_run = heading.add_run(
        f"{number}. "
    )
    number_run.bold = True
    number_run.font.size = Pt(12)
    _set_run_font(number_run)

    if heading_url:
        _add_hyperlink(
            heading,
            item["title"],
            heading_url,
            bold=True,
            size=12,
        )
    else:
        title_run = heading.add_run(
            item["title"]
        )
        title_run.bold = True
        title_run.font.size = Pt(12)
        _set_run_font(title_run)

    if item["type"] == "signal":
        meta = document.add_paragraph()

        meta.paragraph_format.space_after = (
            Pt(5)
        )

        meta_run = meta.add_run(
            "Сигнал МІП: "
            f"{item.get('source_count') or 0} джерел"
            " · "
            f"{item.get('occurrence_count') or 0} публікацій"
        )

        meta_run.italic = True
        meta_run.font.size = Pt(10)
        _set_run_font(meta_run)

        signal_summary = (
            item.get("signal_summary")
            or ""
        ).strip()

        signal_theses = (
            item.get("signal_theses")
            or []
        )

        if signal_summary:
            _add_block_label(
                document,
                "Коротко",
            )

            paragraph = document.add_paragraph(
                signal_summary
            )
            _normalize_paragraph_font(
                paragraph,
                size=12,
            )
            paragraph.paragraph_format.space_after = (
                Pt(6)
            )

        if signal_theses:
            _add_block_label(
                document,
                "Основні тези",
            )

            for index, thesis in enumerate(
                signal_theses,
                start=1,
            ):
                paragraph = (
                    document.add_paragraph()
                )

                paragraph.paragraph_format.left_indent = (
                    Mm(5)
                )
                paragraph.paragraph_format.space_after = (
                    Pt(3)
                )

                number_part = paragraph.add_run(
                    f"{index}) "
                )
                number_part.bold = True
                number_part.font.size = Pt(11)
                _set_run_font(
                    number_part
                )

                thesis_run = paragraph.add_run(
                    thesis
                )
                thesis_run.font.size = Pt(11)
                _set_run_font(
                    thesis_run
                )

        if isinstance(
            example,
            dict,
        ):
            _add_example_publication(
                document,
                example,
            )

        elif item["factual_summary"]:
            paragraph = (
                document.add_paragraph(
                    item[
                        "factual_summary"
                    ]
                )
            )
            _normalize_paragraph_font(
                paragraph,
                size=12,
            )

    else:
        if item["factual_summary"]:
            paragraph = (
                document.add_paragraph(
                    item[
                        "factual_summary"
                    ]
                )
            )
            _normalize_paragraph_font(
                paragraph,
                size=12,
            )
            paragraph.paragraph_format.space_after = (
                Pt(5)
            )

        if (
            item.get("source_name")
            or item.get("source_url")
        ):
            _add_source_line(
                document,
                source_name=(
                    item.get(
                        "source_name"
                    )
                    or ""
                ),
                source_url=(
                    item.get(
                        "source_url"
                    )
                    or ""
                ),
                observed_at=(
                    item.get(
                        "observed_at"
                    )
                ),
            )

    if item["analyst_assessment"]:
        _add_block_label(
            document,
            "Оцінка аналітика",
        )

        paragraph = document.add_paragraph(
            item[
                "analyst_assessment"
            ]
        )
        _normalize_paragraph_font(
            paragraph,
            size=12,
        )
        paragraph.paragraph_format.space_after = (
            Pt(8)
        )


def _add_block_label(
    document: Document,
    text: str,
) -> None:
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = (
        Pt(4)
    )
    paragraph.paragraph_format.space_after = (
        Pt(2)
    )
    paragraph.paragraph_format.keep_with_next = True

    run = paragraph.add_run(
        text
    )
    run.bold = True
    run.italic = True
    run.font.size = Pt(11)
    _set_run_font(run)


def _add_example_publication(
    document: Document,
    example: dict,
) -> None:
    _add_block_label(
        document,
        "Приклад публікації",
    )

    _add_source_line(
        document,
        source_name=(
            example.get(
                "source_name"
            )
            or ""
        ),
        source_url=(
            example.get("url")
            or ""
        ),
        observed_at=(
            example.get(
                "observed_at"
            )
        ),
    )

    text = (
        example.get("text")
        or example.get("title")
        or ""
    ).strip()

    if text:
        paragraph = document.add_paragraph()

        paragraph.paragraph_format.left_indent = (
            Mm(7)
        )
        paragraph.paragraph_format.space_before = (
            Pt(2)
        )
        paragraph.paragraph_format.space_after = (
            Pt(8)
        )

        run = paragraph.add_run(
            text
        )
        run.italic = True
        run.font.size = Pt(11)
        _set_run_font(run)


def _add_source_line(
    document: Document,
    *,
    source_name: str,
    source_url: str,
    observed_at,
) -> None:
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = (
        Pt(3)
    )

    parts = []

    if source_name:
        parts.append(source_name)

    formatted_time = _format_kyiv_time(
        observed_at
    )

    if formatted_time:
        parts.append(formatted_time)

    if parts:
        run = paragraph.add_run(
            " · ".join(parts)
        )
        run.italic = True
        run.font.size = Pt(10)
        _set_run_font(run)

    if source_url:
        if parts:
            separator = paragraph.add_run(
                " · "
            )
            separator.font.size = Pt(10)
            _set_run_font(separator)

        _add_hyperlink(
            paragraph,
            "Відкрити публікацію ↗",
            source_url,
            size=10,
        )


def _format_kyiv_time(
    value,
) -> str:
    if not value:
        return ""

    try:
        parsed = (
            value
            if isinstance(
                value,
                datetime,
            )
            else datetime.fromisoformat(
                str(value)
            )
        )

        if parsed.tzinfo is None:
            return ""

        return (
            parsed
            .astimezone(
                ZoneInfo(
                    "Europe/Kyiv"
                )
            )
            .strftime(
                "%d.%m.%Y %H:%M"
            )
        )

    except (
        TypeError,
        ValueError,
    ):
        return ""


def _add_hyperlink(
    paragraph,
    text: str,
    url: str,
    *,
    bold: bool = False,
    size: int = 10,
) -> None:
    relationship_id = (
        paragraph.part.relate_to(
            url,
            RT.HYPERLINK,
            is_external=True,
        )
    )

    hyperlink = OxmlElement(
        "w:hyperlink"
    )
    hyperlink.set(
        qn("r:id"),
        relationship_id,
    )

    run = OxmlElement("w:r")
    properties = OxmlElement(
        "w:rPr"
    )

    fonts = OxmlElement(
        "w:rFonts"
    )

    for key in (
        "w:ascii",
        "w:hAnsi",
        "w:eastAsia",
        "w:cs",
    ):
        fonts.set(
            qn(key),
            "Times New Roman",
        )

    properties.append(fonts)

    color = OxmlElement(
        "w:color"
    )
    color.set(
        qn("w:val"),
        "0563C1",
    )
    properties.append(color)

    underline = OxmlElement(
        "w:u"
    )
    underline.set(
        qn("w:val"),
        "single",
    )
    properties.append(
        underline
    )

    if bold:
        bold_element = OxmlElement(
            "w:b"
        )
        properties.append(
            bold_element
        )

    size_element = OxmlElement(
        "w:sz"
    )
    size_element.set(
        qn("w:val"),
        str(size * 2),
    )
    properties.append(
        size_element
    )

    run.append(properties)

    text_element = OxmlElement(
        "w:t"
    )
    text_element.text = text
    run.append(text_element)

    hyperlink.append(run)

    paragraph._p.append(
        hyperlink
    )



def _normalize_paragraph_font(
    paragraph,
    *,
    size: int,
) -> None:
    for run in paragraph.runs:
        run.font.name = (
            "Times New Roman"
        )
        run.font.size = Pt(size)
        _set_run_font(run)


def _set_run_font(run) -> None:
    run.font.name = "Times New Roman"

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()

    for key in (
        "w:ascii",
        "w:hAnsi",
        "w:eastAsia",
        "w:cs",
    ):
        rfonts.set(
            qn(key),
            "Times New Roman",
        )


def _required_text(
    value,
    *,
    field: str,
    max_chars: int,
) -> str:
    result = _optional_text(
        value,
        field=field,
        max_chars=max_chars,
    )

    if not result:
        raise ValueError(
            f"{field}: поле порожнє"
        )

    return result


def _optional_text(
    value,
    *,
    field: str,
    max_chars: int,
) -> str:
    if value is None:
        return ""

    if not isinstance(value, str):
        raise ValueError(
            f"{field}: очікується текст"
        )

    result = value.strip()

    if len(result) > max_chars:
        raise ValueError(
            f"{field}: перевищено ліміт"
        )

    return result


def _optional_nonnegative_int(
    value,
    *,
    field: str,
) -> int | None:
    if value is None:
        return None

    if (
        type(value) is not int
        or value < 0
    ):
        raise ValueError(
            f"{field}: очікується невід'ємне число"
        )

    return value


def _optional_url(
    value,
    *,
    field: str,
) -> str:
    result = _optional_text(
        value,
        field=field,
        max_chars=MAX_SOURCE_URL_CHARS,
    )

    if not result:
        return ""

    try:
        parsed = urlsplit(result)

        _ = parsed.port

        if (
            parsed.scheme.lower()
            not in {"http", "https"}
            or not parsed.hostname
            or parsed.username
            or parsed.password
        ):
            raise ValueError

    except ValueError as exc:
        raise ValueError(
            f"{field}: некоректне посилання"
        ) from exc

    return result
